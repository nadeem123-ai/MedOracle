"""
ReportGenerator — produces downloadable .docx and JSON clinical reports.
Uses python-docx (free). No external API calls.
"""
from __future__ import annotations
import io, json
from datetime import datetime
from typing import Dict, Any

_Document = None

def _get_docx():
    global _Document
    if _Document is None:
        try:
            from docx import Document
            _Document = Document
        except ImportError:
            pass
    return _Document


def generate_docx_report(state: Dict[str, Any]) -> bytes:
    """Generate formatted .docx clinical report."""
    Document = _get_docx()
    if Document is None:
        return b""

    doc = Document()
    doc.add_heading("MedOracle  — Clinical Decision Support Report", level=0)
    doc.add_paragraph(f"Generated: {datetime.now().strftime('%Y-%m-%d %H:%M:%S')}")
    doc.add_paragraph("⚠️ DISCLAIMER: This is AI-generated decision support. Always verify with qualified clinicians.")
    doc.add_paragraph()

    # Triage
    triage = state.get("triage", {})
    doc.add_heading("Triage Assessment", level=1)
    doc.add_paragraph(f"Level: {triage.get('level','')} — {triage.get('label','')}")
    doc.add_paragraph(f"Emergency Type: {triage.get('emergency_type','')}")
    doc.add_paragraph(f"Reasoning: {triage.get('reasoning','')}")
    if triage.get("red_flags"):
        doc.add_paragraph("Red Flags: " + " | ".join(triage["red_flags"]))
    if triage.get("immediate_actions"):
        doc.add_heading("Immediate Actions", level=2)
        for a in triage["immediate_actions"]:
            doc.add_paragraph(f"• {a}", style="List Bullet")

    # Diagnoses
    doc.add_heading("Differential Diagnoses", level=1)
    for d in state.get("diagnoses", []):
        doc.add_heading(f"#{d.get('rank','')} {d.get('name','')} [{d.get('icd10','')}]", level=2)
        doc.add_paragraph(f"Probability: {d.get('probability','')} ({d.get('confidence_pct','')}%)")
        doc.add_paragraph(f"Reasoning: {d.get('reasoning','')}")
        if d.get("supporting_tests"):
            doc.add_paragraph("Tests: " + ", ".join(d["supporting_tests"]))

    # Treatment
    tx = state.get("treatment", {})
    doc.add_heading("Treatment Plan", level=1)
    doc.add_paragraph(f"Evidence Grade: {tx.get('evidence_grade','')} — {tx.get('grade_explanation','')}")
    for section, label in [("immediate","Immediate"), ("short_term","Short-Term"), ("long_term","Long-Term")]:
        items = tx.get(section, [])
        if items:
            doc.add_heading(label, level=2)
            for item in items:
                doc.add_paragraph(f"• {item}", style="List Bullet")

    meds = tx.get("medications_recommended", [])
    if meds:
        doc.add_heading("Recommended Medications", level=2)
        for m in meds:
            doc.add_paragraph(f"{m.get('drug','')} — {m.get('dose','')} {m.get('frequency','')} × {m.get('duration','')}")
            doc.add_paragraph(f"  Rationale: {m.get('rationale','')}")

    if tx.get("contraindications"):
        doc.add_heading("Contraindications", level=2)
        for c in tx["contraindications"]:
            doc.add_paragraph(f"✗ {c}")

    # Drug Interactions
    if state.get("drug_interactions"):
        doc.add_heading("Drug Interactions", level=1)
        for ddi in state["drug_interactions"]:
            doc.add_heading(f"[{ddi.get('severity','')}] {ddi.get('drug_a','')} × {ddi.get('drug_b','')}", level=2)
            doc.add_paragraph(f"Mechanism: {ddi.get('mechanism','')}")
            doc.add_paragraph(f"Clinical Effect: {ddi.get('clinical_effect','')}")
            doc.add_paragraph(f"Management: {ddi.get('management','')}")

    # PubMed Citations
    if state.get("pubmed_citations"):
        doc.add_heading("PubMed Evidence Citations", level=1)
        for pub in state["pubmed_citations"]:
            p = doc.add_paragraph()
            p.add_run(pub.get("title","")).bold = True
            doc.add_paragraph(f"{pub.get('authors','')} · {pub.get('journal','')} ({pub.get('year','')})")
            doc.add_paragraph(pub.get("url",""))

    # Processing times
    doc.add_heading("Processing Times", level=1)
    for agent, ms in state.get("processing_time_ms", {}).items():
        doc.add_paragraph(f"{agent}: {ms}ms")

    buf = io.BytesIO()
    doc.save(buf)
    buf.seek(0)
    return buf.read()


def generate_json_report(state: Dict[str, Any]) -> bytes:
    """Generate structured JSON report."""
    report = {k: v for k, v in state.items()}
    report["generated_at"] = datetime.now().isoformat()
    return json.dumps(report, indent=2, ensure_ascii=False).encode("utf-8")
